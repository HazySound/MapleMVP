"""JS에 노출되는 API와 화면 상태 조립."""
from __future__ import annotations

import base64
import ctypes
import io
import json
import logging
import os
import threading
import time
from datetime import date, datetime, timedelta

import webview

from . import cache, demo, ocr, paths
from .mvp import KST, WINDOW, today_kst, week_start
from .rows import merge_months, normalize_usage
from .scraper import (HOME, MAPLE_HOST, USAGE, USAGE_HOST, NeedsLogin, Scraper,
                      ScrapeError)

log = logging.getLogger(__name__)

# 이월 잔액을 재현할 갱신 횟수. 이 기간 이전의 이월은 0으로 가정한다.
HISTORY_WEEKS = 52

# 보관용 과거 수집: 빈 달이 이만큼 이어지면 멈추고, 아무리 멀어도 이 연수까지만 거슬러 올라간다.
ARCHIVE_EMPTY_STOP = 6
ARCHIVE_YEARS = 10


# 창 최소 크기 (논리 픽셀)
MIN_SIZE = (720, 560)


class _POINT(ctypes.Structure):
    _fields_ = [("x", ctypes.c_long), ("y", ctypes.c_long)]


class _RECT(ctypes.Structure):
    _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long), ("right", ctypes.c_long), ("bottom", ctypes.c_long)]


def _find_window(pid: int, title: str) -> int:
    """이 프로세스에서 제목이 title인 최상위 창 핸들."""
    user32 = ctypes.windll.user32
    found = ctypes.c_void_p(0)
    proc = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)

    def check(hwnd, _):
        owner = ctypes.c_ulong()
        user32.GetWindowThreadProcessId(ctypes.c_void_p(hwnd), ctypes.byref(owner))
        if owner.value == pid:
            buf = ctypes.create_unicode_buffer(256)
            user32.GetWindowTextW(ctypes.c_void_p(hwnd), buf, 256)
            if buf.value == title:
                found.value = hwnd
                return False
        return True

    user32.EnumWindows(proc(check), 0)
    return found.value or 0


def _months(first: date, last: date) -> list[tuple[int, int]]:
    out, y, m = [], first.year, first.month
    while (y, m) <= (last.year, last.month):
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _month_final(y: int, m: int, fetched_at: str) -> bool:
    """그 달이 끝난 뒤에 받은 데이터면 더 바뀌지 않는다."""
    nxt = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
    return datetime.fromisoformat(fetched_at).astimezone(KST).date() >= nxt



def _write_xlsx(path: str, rows: list[dict]) -> None:
    """구매내역을 보기 좋은 엑셀 파일로 저장한다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "구매내역"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4A4E5C")
    line = Side(style="thin", color="D8D8DE")
    ws.append(["날짜", "아이템", "금액(원)"])
    for c in ws[1]:
        c.font, c.fill = head, fill
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=line)
    for r in rows:
        y, m, d = (int(x) for x in r["date"].split("-"))
        ws.append([date(y, m, d), r["item"], r["price"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].number_format = "yyyy-mm-dd"
        row[0].alignment = Alignment(horizontal="center")
        row[2].number_format = "#,##0"
    total_at = ws.max_row + 2
    ws.cell(total_at, 2, f"합계 ({len(rows):,}건)").font = Font(bold=True)
    cell = ws.cell(total_at, 3, sum(r["price"] for r in rows))
    cell.font, cell.number_format = Font(bold=True), "#,##0"
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{max(1, ws.max_row - 2)}"
    wb.save(path)


class Api:
    def __init__(self, demo_mode: bool):
        self._demo = demo_mode
        self._main: webview.Window | None = None
        self._scraper = None if demo_mode else Scraper(self._logged_in)
        self._lock = threading.Lock()
        self._cache = {} if demo_mode else cache.load(paths.CACHE)
        self._maximized = False
        self._main_hwnd = 0
        self._resizing = False
        self._usage_error: str | None = None
        # 로그인한 적이 있는지. 기록이 없으면(예전 버전에서 넘어왔으면) 로그인 창 저장소로 판단한다
        seen = cache.load(paths.SETTINGS).get("loggedIn")
        self._fresh = not demo_mode and not (any(paths.WEBVIEW.glob("*")) if seen is None else seen)
        self._logged_out = self._fresh

    def _attach(self, window) -> None:
        self._main = window

    def _rows(self) -> list[dict]:
        if self._demo:
            return demo.rows(today_kst())
        return merge_months(self._cache.get("months", {}), self._cache.get("usage", {}))

    def _pcroom_weeks(self) -> dict[str, int]:
        return {k: int(v) for k, v in cache.load(paths.PCROOM).get("weeks", {}).items()}

    def _build(self, status: str, message: str | None = None) -> dict:
        """화면에 넘길 원본. 13주 계산과 등급 판정은 전부 화면(TS) 쪽에서 한다."""
        synced = self._cache.get("syncedAt") if not self._demo else datetime.now(KST).isoformat()
        floor = (week_start(today_kst()) - timedelta(weeks=HISTORY_WEEKS + WINDOW + 1)).isoformat()
        rows = [r for r in self._rows() if r["date"] >= floor]
        return {"status": status, "message": message, "syncedAt": synced, "demo": self._demo,
                "loggedOut": self._logged_out, "usageError": self._usage_error,
                "rows": rows, "pcroom": self._pcroom_weeks()}

    def _push(self, fn: str, payload) -> None:
        if self._main:
            self._main.evaluate_js(f"window.__mvp && window.__mvp.{fn}({json.dumps(payload, ensure_ascii=False)})")

    # ---- JS API ----
    def get_state(self) -> dict:
        """캐시된 결과를 바로 돌려준다. 캐시가 없으면 status='empty'."""
        if not self._demo and not self._cache.get("months"):
            if self._fresh:   # 로그인한 적이 없으면 수집을 시도해 볼 것도 없이 로그인부터 안내한다
                self._remember_login(False)
                return {"status": "needs_login", "demo": False}
            return {"status": "empty", "demo": False}
        return self._build("cached")

    def refresh(self) -> dict:
        with self._lock:
            if self._demo:
                return self._demo_refresh()
            try:
                self._sync()
            except NeedsLogin:
                log.info("로그인이 필요합니다")
                self._remember_login(False)   # 다음에 열 때는 바로 로그인부터 안내한다
                if self._cache.get("months"):
                    return self._build("needs_login")
                return {"status": "needs_login", "demo": False}
            except ScrapeError as e:
                log.warning("동기화 실패: %s", e)
                if self._cache.get("months"):
                    return self._build("error", str(e))
                return {"status": "error", "message": str(e), "demo": False}
            self._remember_login(True)
            return self._build("ok")

    def hide_login(self) -> None:
        if self._scraper:
            self._scraper.park()

    # ---- PC방 반영액 보정 ----
    def pcroom_save(self, weeks: dict) -> dict:
        """검수를 마친 보정값을 저장한다. 주 시작일이 키라서 주가 지나면 알아서 밀려난다."""
        saved = self._pcroom_weeks()
        saved.update({str(k): max(0, int(v)) for k, v in (weeks or {}).items()})
        floor = (week_start(today_kst()) - timedelta(weeks=HISTORY_WEEKS + WINDOW)).isoformat()
        saved = {k: v for k, v in saved.items() if k >= floor}   # 13주 창을 한참 벗어난 것은 버린다
        cache.save(paths.PCROOM, {"weeks": saved})
        log.info("PC방 보정 저장: %d주, 합계 %s원", len(saved), f"{sum(saved.values()):,}")
        return self._build("ok")

    # ---- 캡처에서 읽기 ----
    def _image(self, data_url: str):
        """data_url이 있으면 그걸, 없으면 클립보드 이미지를 연다."""
        from PIL import Image, ImageGrab
        if data_url:
            try:
                return Image.open(io.BytesIO(base64.b64decode(data_url.partition(",")[2])))
            except Exception as e:
                log.warning("붙여넣은 이미지를 열지 못했어요: %s", e)
                return None
        try:
            got = ImageGrab.grabclipboard()
        except Exception as e:
            log.warning("클립보드를 읽지 못했어요: %s", e)
            return None
        if isinstance(got, list):
            got = Image.open(got[0]) if got else None
        return got if isinstance(got, Image.Image) else None

    def pcroom_scan(self, data_url: str = "", scale: float = 0.0) -> dict:
        """캡처에서 숫자 후보를 뽑아 준다. 어느 것이 맞는지는 화면이 규칙으로 고른다."""
        img = self._image(data_url)
        if img is None:
            return {"ok": False, "message": "이미지가 없어요. 게임 화면에서 PrintScreen을 눌러 주세요."}
        r = ocr.scan(img, float(scale or 0))
        ok = bool(r["readings"] or r["amounts"])
        log.info("캡처 인식: 툴팁 후보 %d, 숫자 후보 %d, 배율 %.2f — %s %s",
                 len(r["readings"]), len(r["amounts"]), r["scale"], img.size, img.mode)
        if not ok:
            try:
                img.save(paths.DATA / "last_capture.png")   # 왜 못 읽었는지 볼 수 있게
            except (OSError, ValueError) as e:
                log.warning("캡처를 남기지 못했어요: %s", e)
        return {"ok": ok, **r,
                "message": "" if ok else "MVP 등급 툴팁을 찾지 못했어요. 등급 게이지에 마우스를 올린 채로 찍어 주세요."}

    def pcroom_clear(self) -> dict:
        cache.save(paths.PCROOM, {"weeks": {}})
        return self._build("ok")

    # ---- 구매내역 보관함 ----
    def history(self, page: int = 1, size: int = 50, q: str = "", start: str = "", end: str = "",
                sort: str = "date", desc: bool = True) -> dict:
        rows = self._filter(q, start, end, sort, bool(desc))
        size = max(1, int(size))
        pages = max(1, -(-len(rows) // size))
        page = min(max(1, int(page)), pages)
        every = self._rows()
        return {
            "rows": rows[(page - 1) * size: page * size],
            "total": len(rows),
            "sum": sum(r["price"] for r in rows),
            "page": page,
            "pages": pages,
            "allTotal": len(every),
            "first": min((r["date"] for r in every), default=""),
            "last": max((r["date"] for r in every), default=""),
            "archived": bool(self._cache.get("archiveDone")),
        }

    def export_history(self, q: str = "", start: str = "", end: str = "",
                       sort: str = "date", desc: bool = True) -> dict:
        # 지금 보고 있는 목록을 엑셀 파일로 저장한다
        rows = self._filter(q, start, end, sort, bool(desc))
        chosen = self._main.create_file_dialog(
            webview.FileDialog.SAVE,
            save_filename=f"메이플_구매내역_{today_kst().isoformat()}.xlsx",
            file_types=("Excel 통합 문서 (*.xlsx)",),
        )
        if not chosen:
            return {"canceled": True}
        path = chosen if isinstance(chosen, str) else chosen[0]
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            _write_xlsx(path, rows)
        except OSError as e:
            log.warning("xlsx export failed: %s", e)
            return {"error": e.strerror or str(e)}
        return {"path": path, "name": os.path.basename(path), "count": len(rows)}

    def _filter(self, q: str, start: str, end: str, sort: str, desc: bool) -> list[dict]:
        text = (q or "").strip().lower()
        rows = [r for r in self._rows()
                if (not start or r["date"] >= start)
                and (not end or r["date"] <= end)
                and (not text or text in r["item"].lower())]
        keys = {
            "date": lambda r: (r["date"], r["item"]),
            "item": lambda r: (r["item"], r["date"]),
            "price": lambda r: (r["price"], r["date"]),
        }
        rows.sort(key=keys.get(sort, keys["date"]), reverse=desc)
        return rows

    def get_ui(self) -> dict:
        return cache.load(paths.UI)

    def save_ui(self, data: dict) -> None:
        cache.save(paths.UI, data)

    def get_plan(self) -> dict:
        return cache.load(paths.PLAN)

    def save_plan(self, data: dict) -> None:
        cache.save(paths.PLAN, data)

    def open_login(self) -> None:
        self._scraper.open_login()

    def minimize(self) -> None:
        self._main.minimize()

    def toggle_maximize(self) -> bool:
        """최대화/복원을 토글하고 바뀐 상태를 돌려준다 (창 단추 아이콘을 바꿔야 해서)."""
        # pywebview에 최대화 상태 조회가 없어서 직접 기억한다
        if self._maximized:
            self._main.restore()
        else:
            self._main.maximize()
        self._maximized = not self._maximized
        return self._maximized

    def is_maximized(self) -> bool:
        return self._maximized

    def close(self) -> None:
        self._main.destroy()

    def start_resize(self, edge: str) -> None:
        """프레임 없는 창은 Windows 기본 크기 조절이 안 돼서, 마우스를 뗄 때까지 커서를 따라 직접 크기를 바꾼다."""
        if self._resizing or self._maximized:
            return
        hwnd = self._hwnd()
        if hwnd:
            self._resizing = True
            threading.Thread(target=self._resize_loop, args=(hwnd, edge), daemon=True).start()

    def _resize_loop(self, hwnd: int, edge: str) -> None:
        user32 = ctypes.windll.user32
        h = ctypes.c_void_p(hwnd)
        try:
            p0, r0 = _POINT(), _RECT()
            user32.GetCursorPos(ctypes.byref(p0))
            user32.GetWindowRect(h, ctypes.byref(r0))
            scale = user32.GetDpiForWindow(h) / 96
            min_w, min_h = int(MIN_SIZE[0] * scale), int(MIN_SIZE[1] * scale)
            last = None
            while user32.GetAsyncKeyState(0x01) & 0x8000:  # 왼쪽 버튼을 누르고 있는 동안
                p = _POINT()
                user32.GetCursorPos(ctypes.byref(p))
                dx, dy = p.x - p0.x, p.y - p0.y
                l, t, r, b = r0.left, r0.top, r0.right, r0.bottom
                if "l" in edge:
                    l = min(r0.left + dx, r0.right - min_w)
                if "r" in edge:
                    r = max(r0.right + dx, r0.left + min_w)
                if "t" in edge:
                    t = min(r0.top + dy, r0.bottom - min_h)
                if "b" in edge:
                    b = max(r0.bottom + dy, r0.top + min_h)
                if (l, t, r, b) != last:
                    # SWP_NOZORDER | SWP_NOACTIVATE
                    user32.SetWindowPos(h, None, l, t, r - l, b - t, 0x0004 | 0x0010)
                    last = (l, t, r, b)
                time.sleep(0.012)
        finally:
            self._resizing = False

    def _hwnd(self) -> int:
        if not self._main_hwnd:
            self._main_hwnd = _find_window(os.getpid(), self._main.title)
        return self._main_hwnd

    # ---- 내부 ----
    def _fetch_month(self, y: int, m: int, label: str, done: int, total: int) -> list[dict]:
        self._push("onProgress", {"label": label, "done": done, "total": total, "count": len(self._rows())})
        rows = self._scraper.fetch_month(y, m)
        self._cache["months"][f"{y:04d}-{m:02d}"] = {"rows": rows, "fetchedAt": datetime.now(KST).isoformat()}
        return rows

    def _sync(self) -> None:
        today = today_kst()
        first = week_start(today) - timedelta(weeks=HISTORY_WEEKS + WINDOW)
        months = self._cache.setdefault("months", {})
        todo = [(y, m) for y, m in _months(first, today)
                if not (f"{y:04d}-{m:02d}" in months and _month_final(y, m, months[f"{y:04d}-{m:02d}"]["fetchedAt"]))]
        for i, (y, m) in enumerate(reversed(todo)):  # 최신 달부터: 로그인 문제를 빨리 발견
            self._fetch_month(y, m, f"{y}년 {m}월", i, len(todo))
        try:
            self._sync_usage(first, today)
        except NeedsLogin:
            raise
        except Exception as e:   # 넥슨이 API를 바꿔도 메이플 구매내역만으로 계속 돌아가게 한다
            log.exception("넥슨 통합 사용내역 수집 실패 — 메이플 구매내역만으로 계산합니다")
            self._usage_error = str(e)
        else:
            self._usage_error = None
        self._cache["syncedAt"] = datetime.now(KST).isoformat()
        cache.save(paths.CACHE, self._cache)
        self._archive(today)

    def _sync_usage(self, first: date, today: date) -> None:
        """넥슨 통합 사용내역을 받아 둔다.

        메이플 아이템 구매내역에 안 뜨는 결제가 있어서(예: 게임 내 결제 일부) 이쪽이 정본이다.
        다른 오리진이라 숨김 창을 잠시 옮겼다가 돌려놓는다.
        """
        usage = self._cache.setdefault("usage", {})
        todo = [(y, m) for y, m in _months(first, today)
                if not (f"{y:04d}-{m:02d}" in usage and _month_final(y, m, usage[f"{y:04d}-{m:02d}"]["fetchedAt"]))]
        if not todo:
            return
        self._scraper.goto(USAGE, USAGE_HOST)
        try:
            skipped: dict[str, int] = {}
            for i, (y, m) in enumerate(reversed(todo)):
                self._push("onProgress", {"label": f"넥슨 사용내역 · {y}년 {m}월", "done": i,
                                          "total": len(todo), "count": len(self._rows())})
                rows, skip = normalize_usage(self._scraper.fetch_usage_month(y, m))
                log.info("  %d-%02d 메이플 %d건 %s원", y, m, len(rows),
                         f"{sum(r['price'] for r in rows):,}")
                for k, v in skip.items():
                    skipped[k] = skipped.get(k, 0) + v
                usage[f"{y:04d}-{m:02d}"] = {"rows": rows, "fetchedAt": datetime.now(KST).isoformat()}
            total = sum(len(usage[k]["rows"]) for k in usage)
            log.info("통합 사용내역 %d개월 수집 완료 — 메이플 %d건%s",
                     len(todo), total, f", 제외 {skipped}" if skipped else "")
        finally:
            try:
                self._scraper.goto(HOME, MAPLE_HOST)
            except (NeedsLogin, ScrapeError) as e:
                log.warning("메이플 페이지로 돌아가지 못했어요: %s", e)

    def _archive(self, today: date) -> None:
        """보관용으로 과거 내역을 끝까지 받아 둔다. 한 번 끝을 확인하면 다시 하지 않는다."""
        if self._cache.get("archiveDone"):
            return
        months = self._cache["months"]
        y, m = min(tuple(map(int, k.split("-"))) for k in months)
        floor = (today.year - ARCHIVE_YEARS, today.month)
        empty_streak = 0
        while (y, m) > floor and empty_streak < ARCHIVE_EMPTY_STOP:
            y, m = (y - 1, 12) if m == 1 else (y, m - 1)
            rows = self._fetch_month(y, m, f"지난 내역 보관 · {y}년 {m}월", 0, 0)
            cache.save(paths.CACHE, self._cache)  # 중간에 닫아도 다음에 이어받게
            empty_streak = empty_streak + 1 if not rows else 0
        self._cache["archiveDone"] = True
        cache.save(paths.CACHE, self._cache)

    def _demo_refresh(self) -> dict:
        labels = ["로그인 세션 확인", "이번 달", "지난달", "13주 계산"]
        for i, label in enumerate(labels):
            self._push("onProgress", {"label": label, "done": i, "total": len(labels)})
            time.sleep(0.45)
        return self._build("ok")

    def _logged_in(self) -> None:
        self._remember_login(True)
        self._push("onLoggedIn", {})

    def _remember_login(self, ok: bool) -> None:
        self._logged_out = not ok
        settings = cache.load(paths.SETTINGS)
        settings["loggedIn"] = ok
        cache.save(paths.SETTINGS, settings)

    def _shutdown(self) -> None:
        if self._scraper:
            self._scraper.destroy()
